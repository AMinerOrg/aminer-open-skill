import copy
import importlib.util
import json
import sys
import tempfile
import unittest
from unittest import mock
from pathlib import Path

from openpyxl import load_workbook

SKILL_ROOT = Path(__file__).resolve().parents[1]
MODULE_PATH = SKILL_ROOT / "scripts" / "build_outputs.py"
SPEC = importlib.util.spec_from_file_location("sota_finder_build_outputs", MODULE_PATH)
MODULE = importlib.util.module_from_spec(SPEC)
assert SPEC.loader is not None
sys.modules[SPEC.name] = MODULE
SPEC.loader.exec_module(MODULE)
FIXTURE = SKILL_ROOT / "tests" / "fixtures" / "sample_research.json"


class BuildOutputsTests(unittest.TestCase):
    def raw_fixture(self):
        return json.loads(FIXTURE.read_text(encoding="utf-8"))

    def test_normalization_is_stable_and_respects_direction_and_ties(self):
        first = MODULE.normalize_bundle(self.raw_fixture())
        second = MODULE.normalize_bundle(self.raw_fixture())
        self.assertEqual(
            [(claim["claim_id"], claim["scope_id"], claim["rank"]) for claim in first["claims"]],
            [(claim["claim_id"], claim["scope_id"], claim["rank"]) for claim in second["claims"]],
        )
        quality = next(scope for scope in first["scopes"] if scope["name"] == "Standard quality track")
        quality_main = [claim for claim in first["claims"] if claim["scope_id"] == quality["scope_id"] and claim["section"] == "main"]
        self.assertEqual([claim["rank"] for claim in quality_main], [1, 1])
        self.assertTrue(all(claim["sota"] for claim in quality_main))

        latency = next(scope for scope in first["scopes"] if scope["name"] == "Standard latency track")
        latency_main = [claim for claim in first["claims"] if claim["scope_id"] == latency["scope_id"] and claim["section"] == "main"]
        self.assertEqual([claim["method"] for claim in latency_main], ["Method B", "Method A"])
        self.assertEqual([claim["rank"] for claim in latency_main], [1, 2])

        reference = [claim for claim in first["claims"] if claim["section"] == "reference"]
        self.assertTrue(reference)
        self.assertTrue(all(claim["rank"] is None and not claim["sota"] for claim in reference))
        method_a = next(claim for claim in first["claims"] if claim["method"] == "Method A" and claim["scope_id"] == quality["scope_id"])
        self.assertTrue(method_a["evidence_locators"][0]["source_id"].startswith("src-"))
        self.assertNotIn("source_key", method_a["evidence_locators"][0])

    def test_run_id_shorthand_and_scope_leader_labels_are_normalized(self):
        raw = self.raw_fixture()
        raw["research_run"]["run_id"] = "followup-20261009"
        data = MODULE.normalize_bundle(raw)

        self.assertEqual(data["research_run"]["run_id"], "run-followup-20261009")
        self.assertEqual(MODULE._labels("en")["sota"], "Scope leader")
        self.assertEqual(MODULE._labels("zh")["sota"], "范围第一")
        self.assertEqual(MODULE._labels("en")["methods"], "Leading Method Notes")
        self.assertEqual(MODULE._labels("zh")["methods"], "领先方法简介")

        raw["research_run"]["run_id"] = "run-already-normalized"
        data = MODULE.normalize_bundle(raw)
        self.assertEqual(data["research_run"]["run_id"], "run-already-normalized")

    def test_builds_exact_four_artifacts_and_multiscope_workbook(self):
        data = MODULE.normalize_bundle(self.raw_fixture())
        with tempfile.TemporaryDirectory() as tmp:
            output = MODULE.build_artifacts(data, output_root=Path(tmp))
            self.assertEqual({path.name for path in output.iterdir()}, set(MODULE.DEFAULT_ARTIFACTS))
            evidence = json.loads((output / "evidence.json").read_text(encoding="utf-8"))
            self.assertEqual(evidence["schema_version"], MODULE.SCHEMA_VERSION)
            self.assertEqual(len(evidence["scopes"]), 2)
            self.assertTrue(all(claim["claim_id"].startswith("clm-") for claim in evidence["claims"]))

            workbook = load_workbook(output / "leaderboards.xlsx", read_only=False, data_only=True)
            self.assertEqual(workbook.sheetnames[0], "Index")
            self.assertIn("Standard quality track", workbook.sheetnames)
            self.assertIn("Standard latency track", workbook.sheetnames)
            quality = workbook["Standard quality track"]
            headers = [cell.value for cell in quality[1]]
            self.assertIn("Claim ID", headers)
            self.assertIn("Evidence Source IDs", headers)
            self.assertEqual(quality.freeze_panes, "A2")

            html_text = (output / "full-report.html").read_text(encoding="utf-8")
            self.assertNotIn("cdn.", html_text.lower())
            self.assertIn("Method A", html_text)
            self.assertIn("claim ID", html_text)

    def test_optional_coverage_matrix_is_rendered_in_all_artifacts(self):
        raw = self.raw_fixture()
        raw["coverage"]["matrix"] = {
            "title": "Cross-benchmark coverage",
            "columns": ["Method", "Benchmark A", "Benchmark B"],
            "rows": [
                ["Method A", "92.4", "not evaluated"],
                ["Method B", "91.0", "88.2"],
            ],
            "notes": ["Cells preserve source-native score wording."],
        }
        data = MODULE.normalize_bundle(raw)
        self.assertEqual(data["coverage"]["matrix"]["rows"][0][2], "not evaluated")

        with tempfile.TemporaryDirectory() as tmp:
            output = MODULE.build_artifacts(data, output_root=Path(tmp))
            evidence = json.loads((output / "evidence.json").read_text(encoding="utf-8"))
            self.assertEqual(evidence["coverage"]["matrix"]["columns"][1], "Benchmark A")

            brief = (output / "brief-report.md").read_text(encoding="utf-8")
            html = (output / "full-report.html").read_text(encoding="utf-8")
            self.assertIn("Cross-benchmark coverage", brief)
            self.assertIn("not evaluated", brief)
            self.assertIn("Cross-benchmark coverage", html)
            self.assertIn("not evaluated", html)

            workbook = load_workbook(output / "leaderboards.xlsx", read_only=False, data_only=True)
            self.assertIn("Cross-benchmark coverage", workbook.sheetnames)
            matrix = workbook["Cross-benchmark coverage"]
            self.assertEqual([cell.value for cell in matrix[1]], ["Method", "Benchmark A", "Benchmark B"])
            self.assertEqual(matrix[2][2].value, "not evaluated")

    def test_brief_report_leads_with_explicit_conclusion_and_keeps_short_board_complete(self):
        raw = self.raw_fixture()
        raw["research_run"]["conclusions"] = "Method A is the practical choice for quality."
        quality_key = raw["scopes"][0]["key"]
        for index in range(3, 9):
            claim = copy.deepcopy(raw["claims"][0])
            claim["key"] = f"extra-quality-{index}"
            claim["scope_key"] = quality_key
            claim["method"] = f"Extra Method {index}"
            claim["score_numeric"] = 91.2 - index
            claim["score_display"] = str(91.2 - index)
            claim["evidence_locators"] = []
            raw["claims"].append(claim)
        data = MODULE.normalize_bundle(raw)
        brief = MODULE.render_brief_markdown(data)
        full = MODULE.render_full_html(data)

        self.assertLess(brief.index("## Conclusion"), brief.index("## Leaderboard Overview"))
        self.assertIn("Method A is the practical choice for quality.", brief)
        self.assertNotIn("Evidence Source IDs", brief)
        self.assertNotIn("Claim ID", brief)
        self.assertIn("Extra Method 8", brief)
        self.assertIn("Extra Method 8", full)

    def test_brief_report_truncates_a_genuinely_long_board(self):
        raw = self.raw_fixture()
        quality_key = raw["scopes"][0]["key"]
        for index in range(3, 23):
            claim = copy.deepcopy(raw["claims"][0])
            claim["key"] = f"extra-quality-{index}"
            claim["scope_key"] = quality_key
            claim["method"] = f"Extra Method {index}"
            claim["score_numeric"] = 91.2 - index
            claim["score_display"] = str(91.2 - index)
            claim["evidence_locators"] = []
            raw["claims"].append(claim)
        data = MODULE.normalize_bundle(raw)
        brief = MODULE.render_brief_markdown(data)
        full = MODULE.render_full_html(data)

        self.assertNotIn("Extra Method 22", brief)
        self.assertIn("Extra Method 22", full)
        self.assertIn("See the [full report](full-report.html) and [workbook](leaderboards.xlsx)", brief)

    def test_many_single_row_scopes_use_one_overview_not_repeated_sections(self):
        raw = self.raw_fixture()
        raw["scopes"] = []
        raw["claims"] = []
        for index in range(1, 14):
            raw["scopes"].append({
                "key": f"track-{index}",
                "name": f"Track {index}",
                "task": "classification",
                "benchmark": "SyntheticBench",
                "dataset_or_subtask": f"track-{index}",
                "split": "test",
                "metric": {"name": "accuracy", "unit": "%", "direction": "higher"},
            })
            raw["claims"].append({
                "key": f"leader-{index}",
                "scope_key": f"track-{index}",
                "section": "main",
                "method": f"Leader {index}",
                "score_numeric": 100 - index,
                "score_display": str(100 - index),
                "evidence_source_keys": ["official-board"],
                "comparability_complete": True,
            })
        raw["conflicts"] = []
        data = MODULE.normalize_bundle(raw)
        brief = MODULE.render_brief_markdown(data)

        self.assertIn("Track 13", brief)
        self.assertNotIn("## Track 1", brief)
        self.assertEqual(brief.count("## Leaderboard Overview"), 1)
        self.assertNotIn("See the [full report](full-report.html)", brief)

    def test_collapsed_scope_details_link_to_complete_results(self):
        raw = self.raw_fixture()
        raw["scopes"] = []
        raw["claims"] = []
        for index in range(1, 14):
            raw["scopes"].append({
                "key": f"track-{index}",
                "name": f"Track {index}",
                "task": "classification",
                "benchmark": "SyntheticBench",
                "dataset_or_subtask": f"track-{index}",
                "split": "test",
                "metric": {"name": "accuracy", "unit": "%", "direction": "higher"},
            })
            for method_index in range(1, 3 if index == 1 else 2):
                raw["claims"].append({
                    "key": f"leader-{index}-{method_index}",
                    "scope_key": f"track-{index}",
                    "section": "main",
                    "method": f"Leader {index}.{method_index}",
                    "score_numeric": 100 - index - method_index / 10,
                    "score_display": str(100 - index - method_index / 10),
                    "evidence_source_keys": ["official-board"],
                    "comparability_complete": True,
                })
        raw["conflicts"] = []
        data = MODULE.normalize_bundle(raw)
        brief = MODULE.render_brief_markdown(data)

        self.assertNotIn("## Track 1", brief)
        self.assertIn("See the [full report](full-report.html) and [workbook](leaderboards.xlsx)", brief)

    def test_coverage_matrix_rejects_mismatched_rows(self):
        raw = self.raw_fixture()
        raw["coverage"]["matrix"] = {
            "columns": ["Method", "Benchmark A"],
            "rows": [["Method A"]],
        }
        with self.assertRaisesRegex(MODULE.ContractError, "same number of cells"):
            MODULE.normalize_bundle(raw)

    def test_existing_bundle_is_not_overwritten(self):
        data = MODULE.normalize_bundle(self.raw_fixture())
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            first = MODULE.build_artifacts(data, output_root=root)
            sentinel = (first / "evidence.json").read_text(encoding="utf-8")
            second = MODULE.build_artifacts(data, output_root=root)
            self.assertNotEqual(first, second)
            self.assertEqual(second.name, f"{first.name}-2")
            self.assertEqual((first / "evidence.json").read_text(encoding="utf-8"), sentinel)

    def test_explicit_empty_output_directory_is_reused_without_overwrite(self):
        data = MODULE.normalize_bundle(self.raw_fixture())
        with tempfile.TemporaryDirectory() as tmp:
            target = Path(tmp) / "custom-run"
            target.mkdir()
            (target / "keep.txt").write_text("preserve", encoding="utf-8")
            output = MODULE.build_artifacts(data, output_dir=target)
            self.assertEqual(output, target)
            self.assertEqual((target / "keep.txt").read_text(encoding="utf-8"), "preserve")
            self.assertTrue(all((target / name).exists() for name in MODULE.DEFAULT_ARTIFACTS))

    def test_cli_help_describes_single_explicit_directory(self):
        parser = MODULE.build_parser()
        action = next(action for action in parser._actions if action.dest == "output_dir")
        self.assertIn("single final run directory", action.help)
        self.assertIn("same run_id", action.help)

    def test_explicit_existing_bundle_is_rejected_instead_of_suffixed(self):
        data = MODULE.normalize_bundle(self.raw_fixture())
        with tempfile.TemporaryDirectory() as tmp:
            target = Path(tmp) / "custom-run"
            first = MODULE.build_artifacts(data, output_dir=target)
            with self.assertRaisesRegex(MODULE.ContractError, "existing artifact directory"):
                MODULE.build_artifacts(data, output_dir=target)
            self.assertEqual(first, target)
            self.assertFalse((target.parent / "custom-run-2").exists())

    def test_explicit_sibling_with_same_run_id_is_rejected(self):
        data = MODULE.normalize_bundle(self.raw_fixture())
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            first = MODULE.build_artifacts(data, output_dir=root / "custom-run")
            with self.assertRaisesRegex(MODULE.ContractError, "same|run_id|already exists"):
                MODULE.build_artifacts(data, output_dir=root / "custom-run-2")
            self.assertEqual(first.name, "custom-run")
            self.assertFalse((root / "custom-run-2").exists())

    def test_artifact_failure_leaves_no_partial_bundle(self):
        data = MODULE.normalize_bundle(self.raw_fixture())
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            with mock.patch.object(MODULE, "render_workbook", side_effect=RuntimeError("synthetic failure")):
                with self.assertRaisesRegex(RuntimeError, "synthetic failure"):
                    MODULE.build_artifacts(data, output_root=root)
            target = root / data["research_run"]["topic_slug"]
            self.assertFalse(target.exists())
            self.assertFalse(any(root.glob(".sota-finder-*")))

    def test_invalid_calendar_date_is_rejected(self):
        raw = self.raw_fixture()
        raw["research_run"]["evidence_checked_through"] = "2026-02-31"
        with self.assertRaisesRegex(MODULE.ContractError, "valid calendar date"):
            MODULE.normalize_bundle(raw)

    def test_main_claim_requires_authoritative_evidence(self):
        raw = self.raw_fixture()
        raw["claims"][0]["evidence_source_keys"] = ["aggregator"]
        with self.assertRaisesRegex(MODULE.ContractError, "requires official, primary, or author evidence"):
            MODULE.normalize_bundle(raw)

    def test_unknown_cross_reference_is_rejected(self):
        raw = self.raw_fixture()
        raw["claims"][0]["evidence_source_keys"] = ["missing-source"]
        with self.assertRaisesRegex(MODULE.ContractError, "Unknown source reference"):
            MODULE.normalize_bundle(raw)

    def test_normalization_does_not_add_optional_research_control(self):
        data = MODULE.normalize_bundle(self.raw_fixture())
        self.assertNotIn("research_control", data)

    def test_common_agent_shorthand_is_normalized_without_retry(self):
        raw = self.raw_fixture()
        raw["coverage"].pop("stop_reason", None)
        raw["claims"][0]["setting"] = "standard evaluation"
        raw["claims"][0]["evidence_locators"] = ["Table 2"]

        data = MODULE.normalize_bundle(raw)

        self.assertTrue(data["coverage"]["stop_reason"])
        self.assertEqual(data["claims"][0]["setting"], {"description": "standard evaluation"})
        self.assertEqual(data["claims"][0]["evidence_locators"], [{"description": "Table 2"}])

    def test_model_cannot_self_assert_query_telemetry(self):
        raw = self.raw_fixture()
        raw["research_control"] = {
            "profile": "deep_review",
            "telemetry_available": True,
            "query_families": [
                {
                    "family": "aliases",
                    "requested_terms": ["DocVQA"],
                    "observed_queries": ["DocVQA benchmark leaderboard"],
                    "status": "observed",
                }
            ],
        }
        data = MODULE.normalize_bundle(raw)
        control = data["research_control"]
        self.assertFalse(control["telemetry_available"])
        self.assertEqual(control["query_families"][0]["observed_queries"], [])
        self.assertEqual(control["query_families"][0]["status"], "not_auditable")

    def test_chinese_visible_reports_keep_machine_ids(self):
        raw = self.raw_fixture()
        raw["research_run"]["language"] = "zh-CN"
        raw["research_control"] = {
            "profile": "standard_research",
            "telemetry_available": False,
            "query_families": [
                {
                    "id": "zh-aliases",
                    "requested_terms": ["中文别名"],
                    "observed_queries": ["中文别名 benchmark"],
                    "status": "observed",
                }
            ],
        }
        data = MODULE.normalize_bundle(raw)
        self.assertEqual(data["research_control"]["profile"], "standard_research")
        self.assertEqual(data["research_control"]["query_families"][0]["observed_queries"], [])
        self.assertEqual(data["research_control"]["query_families"][0]["status"], "not_auditable")
        with tempfile.TemporaryDirectory() as tmp:
            output = MODULE.build_artifacts(data, output_root=Path(tmp))
            self.assertRegex(output.name, r"^[a-z0-9]+(?:-[a-z0-9]+)*$")
            brief = (output / "brief-report.md").read_text(encoding="utf-8")
            self.assertIn("# SOTA 与排行榜报告", brief)
            self.assertIn("主排行榜", brief)
            self.assertNotIn("声明 ID", brief)
            self.assertLess(brief.index("## 结论"), brief.index("## 榜单概览"))
            self.assertNotIn("Evaluation protocol", brief)
            self.assertNotIn("Material axes", brief)
            evidence = json.loads((output / "evidence.json").read_text(encoding="utf-8"))
            self.assertTrue(evidence["claims"][0]["claim_id"].startswith("clm-"))
            self.assertEqual(evidence["scopes"][0]["dataset_or_subtask"], raw["scopes"][0]["dataset_or_subtask"])
            self.assertEqual(evidence["scopes"][0]["split"], raw["scopes"][0]["split"])
            self.assertEqual(evidence["scopes"][0]["metric"], raw["scopes"][0]["metric"])
            workbook = load_workbook(output / "leaderboards.xlsx", read_only=False, data_only=True)
            self.assertEqual(workbook["索引"]["A1"].value, "范围")
            self.assertEqual(workbook["索引"]["C1"].value, "基准")
            self.assertEqual(workbook["索引"]["K1"].value, "范围 ID")
            data_sheet = workbook[workbook.sheetnames[1]]
            self.assertEqual(data_sheet["L1"].value, "声明 ID")
            self.assertEqual(data_sheet["B2"].value, "主榜")
            self.assertEqual(data_sheet["D2"].value, "单模型")
            self.assertEqual(data_sheet["J2"].value, "是")

    def test_redundant_comparability_status_is_not_emitted(self):
        raw = self.raw_fixture()
        raw["claims"][0]["comparability_status"] = "partial"
        data = MODULE.normalize_bundle(raw)
        self.assertNotIn("comparability_status", data["claims"][0])


if __name__ == "__main__":
    unittest.main()
